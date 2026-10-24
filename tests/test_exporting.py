"""Dashboard exports: Excel workbook and the printable summary."""
import tempfile
import unittest
from pathlib import Path

from backend.analysis import dataset, summarize
from desktop.exporting import kpis, summary_html
from tests.test_service import Base
from tests.test_sms import FakeSender

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


class Exports(Base):
    def setUp(self):
        super().setUp()
        self.svc.sms = FakeSender()
        self.svc.sync_notifications = True
        for i, (loc, issues) in enumerate([("2F Male", ["no_soap"]), ("2F Male", ["no_soap"]), ("GF Female", ["leak"])]):
            self.svc.store.set(self.svc.C_PUBLIC_REPORTS, f"r{i:019d}", {
                "location": loc, "issues": issues, "urgency": "medium", "rating": 3, "deviceId": str(i) * 24,
                "source": "qr"})
        self.svc.import_public_reports()
        self.records = dataset(self.svc)
        self.summary = summarize(self.records)

    def test_kpis_and_html(self):
        self.assertEqual(dict(kpis(self.summary))["Reports"], "3")
        self.assertEqual(dict(kpis(self.summary))["Avg cleanliness"], "3/5")
        page = summary_html(self.summary, 30)
        self.assertIn("Reports by restroom", page)
        self.assertIn("No soap at 2F Male: 2 times", page)

    @unittest.skipIf(openpyxl is None, "openpyxl not installed")
    def test_xlsx(self):
        from desktop.exporting import write_xlsx
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.xlsx"
            write_xlsx(str(path), self.summary, self.records, 30)
            wb = openpyxl.load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Summary", "Reports"])
            self.assertEqual(wb["Reports"].max_row, 4)  # header + 3 reports
            self.assertEqual(wb["Reports"]["A1"].value, "#")


if __name__ == "__main__":
    unittest.main()
