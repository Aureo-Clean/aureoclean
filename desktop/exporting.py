"""Export the reports dashboard to Excel (.xlsx) and PDF.

The Excel and HTML builders don't need Qt, so they can be tested anywhere;
only ``write_pdf`` uses Qt to print the HTML to a PDF.
"""
from __future__ import annotations

import html
from datetime import datetime

CATEGORY_LABELS = {"emergency": "Emergency", "facility": "Broken facility", "supplies": "Supplies",
                   "cleanliness": "Cleanliness", "system": "System problem", "other": "Other"}
COLUMNS = [("ref", "#"), ("createdAt", "Reported"), ("location", "Restroom"), ("category", "Type"),
           ("priority", "Priority"), ("status", "Status"), ("issues", "Problems"), ("urgency", "Urgency"),
           ("rating", "Cleanliness (1-5)"), ("takenBy", "Taken by"), ("minutesToAccept", "Minutes to accept"),
           ("minutesToResolve", "Minutes to fix"), ("escalated", "Escalated"), ("channel", "From"),
           ("details", "Details")]


def kpis(summary: dict) -> list[tuple[str, str]]:
    def mins(v):
        return "—" if v is None else f"{v} min"
    return [
        ("Reports", str(summary.get("totalReports", 0))),
        ("Still open", str(summary.get("openReports", 0))),
        ("Emergencies", str(summary.get("emergencies", 0))),
        ("Median time to accept", mins(summary.get("medianMinutesToAccept"))),
        ("Median time to fix", mins(summary.get("medianMinutesToResolve"))),
        ("Avg cleanliness", "—" if summary.get("averageCleanliness") is None else f"{summary['averageCleanliness']}/5"),
    ]


def _local(ts: str | None) -> str:
    if not ts:
        return ""
    try:
        return datetime.fromisoformat(ts).astimezone().strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return ts


def write_xlsx(path: str, summary: dict, records: list[dict], days: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    navy, soft = "14336E", "E6F0FB"
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "AureoClean report summary"
    ws["A1"].font = Font(bold=True, size=14, color=navy)
    ws["A2"] = f"Last {days} days · exported {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    row = 4
    for name, value in kpis(summary):
        ws.cell(row, 1, name).font = Font(bold=True)
        ws.cell(row, 2, value)
        row += 1

    def table(title, mapping, headers):
        nonlocal row
        row += 1
        ws.cell(row, 1, title).font = Font(bold=True, color=navy, size=12)
        row += 1
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row, c, h)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=soft)
        row += 1
        for k, v in mapping:
            ws.cell(row, 1, k)
            for c, val in enumerate(v if isinstance(v, (list, tuple)) else [v], 2):
                ws.cell(row, c, val)
            row += 1

    table("By restroom", summary.get("byLocation", {}).items(), ["Restroom", "Reports"])
    table("By problem", summary.get("byIssue", {}).items(), ["Problem", "Times"])
    table("By type", [(CATEGORY_LABELS.get(k, k), v) for k, v in summary.get("byCategory", {}).items()],
          ["Type", "Reports"])
    table("By hour of day", [(f"{h:02d}:00", n) for h, n in sorted(
        (int(k), v) for k, v in summary.get("byHour", {}).items())], ["Hour", "Reports"])
    table("Recurring problems", [(p["location"], (p["issue"], p["count"])) for p in summary.get("recurringProblems", [])],
          ["Restroom", "Problem", "Times"])
    table("Janitors", [(name, (v["handled"], v["medianMinutesToResolve"]))
                       for name, v in summary.get("byJanitor", {}).items()],
          ["Janitor", "Jobs handled", "Median minutes to fix"])
    ws.column_dimensions["A"].width = 30
    for col in "BCD":
        ws.column_dimensions[col].width = 18

    data = wb.create_sheet("Reports")
    for c, (_, header) in enumerate(COLUMNS, 1):
        cell = data.cell(1, c, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
    for r, rec in enumerate(records, 2):
        for c, (key, _) in enumerate(COLUMNS, 1):
            value = rec.get(key)
            if key == "createdAt":
                value = _local(value)
            elif key == "issues":
                value = ", ".join(value or [])
            elif key == "category":
                value = CATEGORY_LABELS.get(value, value)
            elif key == "escalated":
                value = "yes" if value else ""
            data.cell(r, c, value)
    widths = {"#": 6, "Reported": 17, "Restroom": 22, "Problems": 34, "Details": 50}
    for c, (_, header) in enumerate(COLUMNS, 1):
        data.column_dimensions[get_column_letter(c)].width = widths.get(header, 14)
    data.freeze_panes = "A2"
    data.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(records) + 1)}"
    for row_cells in data.iter_rows(min_row=2):
        row_cells[-1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def summary_html(summary: dict, days: int) -> str:
    """Printable summary (Qt rich text: simple tables, bars drawn with cell widths)."""
    e = html.escape

    def bars(title, mapping, top=10):
        items = list(mapping)[:top]
        if not items:
            return f"<h3>{e(title)}</h3><p style='color:#5B6B80'>No data.</p>"
        biggest = max(v for _, v in items) or 1
        rows = "".join(
            f"<tr><td width='34%'>{e(str(k))}</td><td width='56%'><table width='{max(2, int(v / biggest * 100))}%' "
            f"cellspacing='0' cellpadding='0'><tr><td bgcolor='#1F63C6'>&nbsp;</td></tr></table></td>"
            f"<td width='10%' align='right'>{v}</td></tr>" for k, v in items)
        return f"<h3>{e(title)}</h3><table width='100%' cellspacing='0' cellpadding='3'>{rows}</table>"

    tiles = "".join(f"<td align='center' bgcolor='#E6F0FB'><b style='font-size:15pt;color:#14336E'>{e(v)}</b>"
                    f"<br><span style='color:#5B6B80'>{e(k)}</span></td>" for k, v in kpis(summary))
    hours = sorted((int(k), v) for k, v in summary.get("byHour", {}).items())
    recurring = "".join(f"<li>{e(p['issue'])} at {e(p['location'])}: {p['count']} times</li>"
                        for p in summary.get("recurringProblems", [])) or "<li>None yet.</li>"
    janitors = "".join(f"<tr><td>{e(n)}</td><td align='right'>{v['handled']}</td>"
                       f"<td align='right'>{'—' if v['medianMinutesToResolve'] is None else v['medianMinutesToResolve']}"
                       f"</td></tr>" for n, v in summary.get("byJanitor", {}).items())
    return f"""<html><body style="font-family:sans-serif; color:#162338">
<h1 style="color:#14336E">AureoClean report summary</h1>
<p style="color:#5B6B80">Last {days} days · exported {datetime.now().strftime('%B %d, %Y %H:%M')}</p>
<table width="100%" cellspacing="6" cellpadding="8"><tr>{tiles}</tr></table>
{bars("Reports by restroom", summary.get("byLocation", {}).items())}
{bars("Most common problems", summary.get("byIssue", {}).items())}
{bars("Reports by type", [(CATEGORY_LABELS.get(k, k), v) for k, v in summary.get("byCategory", {}).items()])}
{bars("Busiest hours", [(f"{h:02d}:00", n) for h, n in hours], top=24)}
<h3>Recurring problems</h3><ul>{recurring}</ul>
<h3>Janitors</h3>
<table width="100%" cellspacing="0" cellpadding="4"><tr bgcolor="#E6F0FB"><td><b>Janitor</b></td>
<td align="right"><b>Jobs handled</b></td><td align="right"><b>Median minutes to fix</b></td></tr>
{janitors or "<tr><td colspan='3'>No jobs taken by SMS yet.</td></tr>"}</table>
</body></html>"""


def write_pdf(path: str, html_text: str) -> None:
    from PySide6.QtCore import QMarginsF, QSizeF
    from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument

    writer = QPdfWriter(path)
    writer.setPageLayout(QPageLayout(QPageSize(QPageSize.PageSizeId.A4), QPageLayout.Orientation.Portrait,
                                     QMarginsF(15, 15, 15, 15), QPageLayout.Unit.Millimeter))
    writer.setResolution(96)
    doc = QTextDocument()
    doc.setHtml(html_text)
    doc.setPageSize(QSizeF(writer.width(), writer.height()))
    doc.print_(writer)
